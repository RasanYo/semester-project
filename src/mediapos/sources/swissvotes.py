"""Swissvotes: the authority for federal popular votes (block A and block D).

One flat file holds every federal ballot object since 1848 with its date,
bilingual titles, policy domains, party voting recommendations and per-canton
results. Licence CC BY 4.0 (Swissvotes, Annee Politique Suisse, Univ. of Bern).

The file has a gap: every numeric per-canton column is empty for vote days from
2024-06-09 onward in the 2026-07-01 release. Without a repair, no object after
2024-03-03 can be scored on the Roestigraben axis, and the whole second half of
2024 silently leaves the candidate pool. `fill_cantonal_gap` closes it from two
other published sources.
"""

from __future__ import annotations

import csv
import json
from datetime import date
from pathlib import Path
from typing import Any

from ..config import (
    CANTONS_DE,
    CANTONS_EXCLUDED,
    CANTONS_FR,
    EXCLUDED_DOMAIN_PREFIX,
    MERGER_DATE,
    MIGRATION_EU_PREFIXES,
    WINDOW_FROM,
    WINDOW_TO,
)
from ..manifest import Manifest, fetch

DATASET_URL = "https://swissvotes.ch/page/dataset/swissvotes_dataset.csv"
CODEBOOK_URL = "https://swissvotes.ch/page/dataset/codebook-de.pdf"

# Per-object BFS workbook; exists for most but not all objects.
STAATSEBENEN_URL = "https://swissvotes.ch/vote/{anr}.00/staatsebenen-de.xlsx"
# Federal Statistical Office open-government results, one file per vote day.
BFS_OGD_URL = (
    "https://ogd-static.voteinfo-app.ch/v1/ogd/"
    "sd-t-17-02-{ymd}-eidgAbstimmung.json"
)

ALL_CANTONS = CANTONS_FR + CANTONS_DE + CANTONS_EXCLUDED
LICENCE = "CC BY 4.0 -- Swissvotes, Annee Politique Suisse, University of Bern"


def _num(raw: str | None) -> float | None:
    """Parse a Swissvotes numeric cell.

    Missing shows up as "", "." or 9999 ("the organisation did not exist");
    thousands are separated by U+2019 in some monetary columns.
    """
    s = (raw or "").strip().replace("’", "")
    if s in ("", ".", "9999"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _parse_date(raw: str) -> date:
    """Swissvotes dates are DD.MM.YYYY, not ISO."""
    d, m, y = raw.split(".")
    return date(int(y), int(m), int(d))


def download(manifest: Manifest | None = None, *, force: bool = False) -> Path:
    path = fetch(
        DATASET_URL,
        "swissvotes/swissvotes_dataset.csv",
        manifest=manifest,
        name="swissvotes_dataset",
        force=force,
        licence=LICENCE,
        role="block A + block D: federal popular votes",
    )
    fetch(
        CODEBOOK_URL,
        "swissvotes/codebook-de.pdf",
        manifest=manifest,
        name="swissvotes_codebook",
        force=force,
        licence=LICENCE,
        role="code definitions (German only -- the -fr URL serves the same bytes)",
    )
    assert path is not None
    return path


def load(path: Path) -> list[dict[str, Any]]:
    """Read every row, annotated with parsed fields. No filtering yet."""
    # The file is semicolon-separated and UTF-8 with a BOM; the default comma
    # delimiter yields one single 874-name column.
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh, delimiter=";"))

    for r in rows:
        r["_anr"] = r["anr"].strip()  # fractional values like "82.1" exist: keep str
        r["_date"] = _parse_date(r["datum"])
        # The classification is spread over three independent slots, and the
        # migration / EU dimension is frequently the second, not the first.
        r["_domains"] = [
            v.strip()
            for i in (1, 2, 3)
            for j in (1, 2, 3)
            if (v := r.get(f"d{i}e{j}", "")) and v.strip() not in (".", "")
        ]
        r["_is_media"] = any(
            d.startswith(EXCLUDED_DOMAIN_PREFIX) for d in r["_domains"]
        )
        r["_is_migration_eu"] = any(
            d.startswith(MIGRATION_EU_PREFIXES) for d in r["_domains"]
        )
        # Follow the centre party across the CVP+BDP -> Die Mitte rename.
        r["CENTRE"] = r["p-mitte"] if r["_date"] >= MERGER_DATE else r["p-cvp"]
        r["_japroz"] = {kt: _num(r.get(f"{kt}-japroz")) for kt in ALL_CANTONS}
        r["_salience"] = _num(r.get("mediares-tot")) or 0.0
    return rows


def in_window(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [r for r in rows if WINDOW_FROM <= r["_date"] <= WINDOW_TO]


def missing_cantonal(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Rows whose per-canton yes-share is absent for any canton we score on."""
    scored = CANTONS_FR + CANTONS_DE
    return [r for r in rows if any(r["_japroz"][kt] is None for kt in scored)]


# --------------------------------------------------------- the 2024 gap ----


def _from_staatsebenen(path: Path) -> dict[str, float]:
    """Read 'Ja in Prozent' per canton out of a BFS workbook."""
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))

    header_i = next(
        i
        for i, row in enumerate(rows)
        if row and any(isinstance(c, str) and "Ja in Prozent" in c for c in row)
    )
    header = [str(c).strip() if c else "" for c in rows[header_i]]
    i_kt = next(
        i for i, c in enumerate(header)
        if c.startswith("Kanton") and "Nummer" not in c
    )
    i_ja = next(i for i, c in enumerate(header) if "Ja in Prozent" in c)

    out: dict[str, float] = {}
    for row in rows[header_i + 1 :]:
        if not row or not isinstance(row[i_kt], str):
            continue
        name, ja = row[i_kt].strip(), row[i_ja]
        kt = CANTON_NAME_TO_CODE.get(name)
        if kt and isinstance(ja, (int, float)):
            out[kt] = float(ja)
    return out


def _from_bfs_ogd(path: Path, anr: str) -> dict[str, float]:
    """Read per-canton yes-share from the BFS open-data file for a vote day.

    The BFS keys ballot objects by `vorlagenId`, which is the Swissvotes `anr`
    times ten.
    """
    payload = json.loads(path.read_text(encoding="utf-8"))
    target = int(float(anr) * 10)

    def walk(node: Any):
        if isinstance(node, dict):
            if node.get("vorlagenId") == target:
                yield node
            for v in node.values():
                yield from walk(v)
        elif isinstance(node, list):
            for v in node:
                yield from walk(v)

    vorlage = next(walk(payload), None)
    if vorlage is None:
        return {}

    out: dict[str, float] = {}
    for kanton in vorlage.get("kantone", []):
        # geoLevelnummer comes back as a string ("1"), not an int.
        try:
            code = BFS_CANTON_ID_TO_CODE.get(int(kanton.get("geoLevelnummer")))
        except (TypeError, ValueError):
            continue
        ja = kanton.get("resultat", {}).get("jaStimmenInProzent")
        if code and ja is not None:
            out[code] = float(ja)
    return out


def fill_cantonal_gap(
    rows: list[dict[str, Any]], manifest: Manifest | None = None
) -> dict[str, str]:
    """Repair `_japroz` in place for rows the main file left empty.

    Returns anr -> which source supplied the numbers, so the manifest can say
    exactly which values did not come from swissvotes_dataset.csv.
    """
    repaired: dict[str, str] = {}
    for r in missing_cantonal(rows):
        anr = r["_anr"]

        wb = fetch(
            STAATSEBENEN_URL.format(anr=anr),
            f"swissvotes/staatsebenen_{anr}.xlsx",
            manifest=manifest,
            name=f"staatsebenen_{anr}",
            allow_missing=True,
            licence=LICENCE,
            role=f"per-canton results for ballot object {anr}",
        )
        values = _from_staatsebenen(wb) if wb else {}
        source = "swissvotes:staatsebenen-de.xlsx"

        if not values:  # the workbook 404s for some objects; BFS still has them
            ymd = r["_date"].strftime("%Y%m%d")
            ogd = fetch(
                BFS_OGD_URL.format(ymd=ymd),
                f"bfs/eidgAbstimmung_{ymd}.json",
                manifest=manifest,
                name=f"bfs_ogd_{ymd}",
                allow_missing=True,
                licence="Open Government Data, Federal Statistical Office",
                role=f"per-canton results for vote day {r['_date']}",
            )
            values = _from_bfs_ogd(ogd, anr) if ogd else {}
            source = "bfs:ogd-static.voteinfo-app.ch"

        if values:
            for kt, ja in values.items():
                if r["_japroz"].get(kt) is None:
                    r["_japroz"][kt] = ja
            repaired[anr] = source
    return repaired


# German canton names as the BFS workbooks spell them.
CANTON_NAME_TO_CODE = {
    "Zürich": "zh", "Bern": "be", "Bern / Berne": "be", "Luzern": "lu",
    "Uri": "ur",
    "Schwyz": "sz", "Obwalden": "ow", "Nidwalden": "nw", "Glarus": "gl", "Zug": "zg",
    "Freiburg": "fr", "Fribourg": "fr", "Freiburg / Fribourg": "fr",
    "Solothurn": "so", "Basel-Stadt": "bs", "Basel-Landschaft": "bl",
    "Schaffhausen": "sh", "Appenzell Ausserrhoden": "ar", "Appenzell Innerrhoden": "ai",
    "St. Gallen": "sg", "Sankt Gallen": "sg", "Graubünden": "gr",
    "Graubünden / Grigioni / Grischun": "gr", "Aargau": "ag", "Thurgau": "tg",
    "Tessin": "ti", "Ticino": "ti", "Waadt": "vd", "Vaud": "vd",
    "Wallis": "vs", "Valais": "vs", "Wallis / Valais": "vs",
    "Neuenburg": "ne", "Neuchâtel": "ne", "Genf": "ge", "Genève": "ge",
    "Jura": "ju",
}

# BFS geoLevelnummer -> canton code (official canton numbering, 1..26).
BFS_CANTON_ID_TO_CODE = {
    1: "zh", 2: "be", 3: "lu", 4: "ur", 5: "sz", 6: "ow", 7: "nw", 8: "gl",
    9: "zg", 10: "fr", 11: "so", 12: "bs", 13: "bl", 14: "sh", 15: "ar",
    16: "ai", 17: "sg", 18: "gr", 19: "ag", 20: "tg", 21: "ti", 22: "vd",
    23: "vs", 24: "ne", 25: "ge", 26: "ju",
}
